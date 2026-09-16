"""Planilha de verdade: numero como numero, data como data (AC-2).

O valor vai tipado para a celula, e o formato de exibicao vai a parte.
Texto que parece formula e' neutralizado como no CSV: o openpyxl grava como
formula qualquer string que comece com `=`.
"""

import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from estoque.application.exportacao.porta import Cabecalho
from estoque.application.exportacao.tabelas import Celula, Coluna, Tabela
from estoque.exportacao._celula import local, texto

FORMATO: dict[str, str] = {
    "inteiro": "#,##0",
    "decimal": "#,##0.0",
    "reais": '"R$" #,##0.00',
    "data": "dd/mm/yyyy",
    "datahora": "dd/mm/yyyy hh:mm",
}


def _valor(v: Celula, c: Coluna) -> str | int | float | date | datetime | None:
    if v is None or isinstance(v, int):
        return v
    if isinstance(v, datetime):
        return local(v)
    if isinstance(v, date):
        return v
    if isinstance(v, Decimal):
        # float na celula: o Excel nao tem decimal exato, e a exibicao arredonda.
        return float(v)
    return texto(v, c)


class ExportadorXlsx:
    extensao = "xlsx"
    tipo_midia = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def gerar(self, tabela: Tabela, cabecalho: Cabecalho) -> bytes:
        wb = Workbook()
        # `active` e' opcional so' no tipo: pasta nova sempre tem uma aba.
        ws = wb.worksheets[0]
        ws.title = texto(tabela.titulo, Coluna(""))[:31]
        ws.append([texto(tabela.titulo, Coluna(""))])
        ws["A1"].font = Font(bold=True, size=13)
        for ctx in tabela.contexto:
            ws.append([texto(ctx, Coluna(""))])
        ws.append(
            [
                f"{cabecalho.emissor} · gerado por {texto(cabecalho.gerado_por, Coluna(''))} "
                f"em {local(cabecalho.gerado_em):%d/%m/%Y %H:%M}"
            ]
        )
        ws.append([])

        inicio = ws.max_row + 1
        ws.append([texto(c.titulo, c) for c in tabela.colunas])
        for titulo in ws[inicio]:
            titulo.font = Font(bold=True)
        for linha in tabela.linhas:
            ws.append([_valor(v, c) for v, c in zip(linha, tabela.colunas, strict=True)])

        for i, c in enumerate(tabela.colunas, start=1):
            letra = get_column_letter(i)
            fmt = FORMATO.get(c.tipo)
            if fmt:
                for r in range(inicio + 1, ws.max_row + 1):
                    ws.cell(row=r, column=i).number_format = fmt
            largura = max([len(c.titulo)] + [len(texto(x[i - 1], c)) for x in tabela.linhas])
            ws.column_dimensions[letra].width = min(max(largura + 2, 10), 60)
        # O cabecalho da tabela fica visivel ao rolar.
        ws.freeze_panes = f"A{inicio + 1}"

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
