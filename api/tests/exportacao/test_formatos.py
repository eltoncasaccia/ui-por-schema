"""T-054 — os tres formatos. AC-2, AC-3, AC-7.

Cada arquivo e' lido de volta como quem o recebe o leria: o CSV pelos bytes,
o XLSX pelo openpyxl, o PDF pelo texto extraido.
"""

import csv
import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from estoque.application.exportacao.porta import Cabecalho, neutralizar
from estoque.application.exportacao.tabelas import Celula, Coluna, Tabela
from estoque.exportacao import EXPORTADORES

AGORA = datetime(2026, 9, 16, 15, 30, tzinfo=UTC)  # 12:30 em Sao Paulo
CAB = Cabecalho("Bertoni Distribuidora Farmacêutica", "Helena Prado", AGORA)
HOSTIL = '=HYPERLINK("http://mal.example","clique")'


def _tabela(linhas: int = 2) -> Tabela:
    base: list[tuple[Celula, ...]] = [
        (
            "Solução fisiológica",
            1200,
            Decimal("1234.5"),
            date(2026, 10, 1),
            datetime(2026, 9, 1, 3, 0, tzinfo=UTC),
        ),
        (HOSTIL, -3, Decimal("0.1"), None, None),
    ]
    return Tabela(
        titulo="Movimentação por unidade",
        colunas=(
            Coluna("Produto"),
            Coluna("Quantidade", "inteiro"),
            Coluna("Valor (R$)", "reais"),
            Coluna("Validade", "data"),
            Coluna("Data", "datahora"),
        ),
        linhas=(base * linhas)[:linhas],
        contexto=["2 unidades", "últimos 90 dias · só efetivados"],
        nome_base="movimentacao",
    )


def _csv_linhas() -> tuple[bytes, list[list[str]]]:
    bruto = EXPORTADORES["csv"].gerar(_tabela(), CAB)
    texto = bruto.decode("utf-8-sig")
    return bruto, list(csv.reader(io.StringIO(texto), delimiter=";"))


def test_ac2_csv_tem_bom_ponto_e_virgula_e_acento_intacto() -> None:
    bruto, linhas = _csv_linhas()
    assert bruto.startswith(b"\xef\xbb\xbf")
    assert "Solução fisiológica".encode() in bruto
    cab = next(i for i, x in enumerate(linhas) if x and x[0] == "Produto")
    assert linhas[cab] == ["Produto", "Quantidade", "Valor (R$)", "Validade", "Data"]
    # Virgula decimal e data brasileira; hora em Sao Paulo (03:00 UTC = 00:00).
    assert linhas[cab + 1] == [
        "Solução fisiológica",
        "1200",
        "1.234,50",
        "01/10/2026",
        "01/09/2026 00:00",
    ]


def test_ac2_csv_traz_o_recorte_e_quem_gerou_antes_da_tabela() -> None:
    _, linhas = _csv_linhas()
    topo = [x[0] for x in linhas[:5] if x]
    assert topo[0] == "Movimentação por unidade"
    assert "últimos 90 dias · só efetivados" in topo
    assert any("Helena Prado" in x and "16/09/2026 12:30" in x for x in topo)


def test_ac2_xlsx_guarda_numero_como_numero_e_data_como_data() -> None:
    wb = load_workbook(io.BytesIO(EXPORTADORES["xlsx"].gerar(_tabela(), CAB)))
    ws = wb.worksheets[0]
    cab = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Produto")
    qtd, valor, validade, quando = (ws.cell(cab + 1, c) for c in (2, 3, 4, 5))
    assert qtd.value == 1200 and qtd.data_type == "n"
    assert valor.value == pytest.approx(1234.5) and "R$" in valor.number_format
    assert isinstance(validade.value, datetime) and validade.value.date() == date(2026, 10, 1)
    assert quando.value == datetime(2026, 9, 1, 0, 0)
    # Negativo continua numero: neutralizar e' so' para texto.
    assert ws.cell(cab + 2, 2).value == -3


def test_ac3_pdf_tem_titulo_recorte_autor_data_e_pagina() -> None:
    bruto = EXPORTADORES["pdf"].gerar(_tabela(), CAB)
    assert bruto.startswith(b"%PDF")
    texto = PdfReader(io.BytesIO(bruto)).pages[0].extract_text()
    for trecho in (
        "BERTONI",
        "Movimentação por unidade",
        "últimos 90 dias",
        "Helena Prado",
        "16/09/2026 12:30",
        "página 1 de 1",
        "1.234,50",
    ):
        assert trecho in texto, trecho


def test_ac3_pdf_longo_numera_todas_as_paginas() -> None:
    bruto = EXPORTADORES["pdf"].gerar(_tabela(linhas=300), CAB)
    paginas = PdfReader(io.BytesIO(bruto)).pages
    assert len(paginas) > 1
    assert f"página {len(paginas)} de {len(paginas)}" in paginas[-1].extract_text()


def test_ac3_pdf_nao_quebra_com_caractere_fora_do_latin1() -> None:
    t = _tabela()
    t.linhas.append(("Ácido — lote “especial” ✓ 药", 1, None, None, None))
    assert EXPORTADORES["pdf"].gerar(t, CAB).startswith(b"%PDF")


def test_pdf_sem_linhas_diz_que_nao_ha_registro() -> None:
    t = Tabela(titulo="Lotes", colunas=(Coluna("Produto"),), linhas=[])
    texto = PdfReader(io.BytesIO(EXPORTADORES["pdf"].gerar(t, CAB))).pages[0].extract_text()
    assert "Nenhum registro" in texto


# --- AC-7 · injecao de formula (negativo) ----------------------------------


@pytest.mark.parametrize("inicio", ["=", "+", "-", "@", "\t", "\r"])
def test_ac7_texto_que_comeca_como_formula_e_neutralizado(inicio: str) -> None:
    assert neutralizar(f"{inicio}1+1") == f"'{inicio}1+1"


def test_ac7_texto_comum_passa_intacto() -> None:
    assert neutralizar("Dipirona 500 mg") == "Dipirona 500 mg"
    assert neutralizar("A-12") == "A-12"


def test_ac7_csv_nao_entrega_formula() -> None:
    _, linhas = _csv_linhas()
    celulas = [c for x in linhas for c in x]
    assert HOSTIL not in celulas
    assert "'" + HOSTIL in celulas


def test_ac7_xlsx_grava_a_formula_como_texto() -> None:
    wb = load_workbook(io.BytesIO(EXPORTADORES["xlsx"].gerar(_tabela(), CAB)))
    ws = wb.worksheets[0]
    celulas = [c for linha in ws.iter_rows() for c in linha if c.value is not None]
    assert all(c.data_type != "f" for c in celulas)
    assert any(c.value == "'" + HOSTIL for c in celulas)


def test_ac7_titulo_e_recorte_tambem_sao_neutralizados() -> None:
    """O recorte vem de dado cadastrado (nome de unidade), nao so' da tabela."""
    t = Tabela(titulo="=1+1", colunas=(Coluna("x"),), linhas=[("a",)], contexto=["@SUM(A1)"])
    wb = load_workbook(io.BytesIO(EXPORTADORES["xlsx"].gerar(t, CAB)))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "'=1+1" and ws["A1"].data_type != "f"
    assert ws["A2"].value == "'@SUM(A1)"
    linhas = EXPORTADORES["csv"].gerar(t, CAB).decode("utf-8-sig").splitlines()
    assert linhas[0] == "'=1+1" and linhas[1] == "'@SUM(A1)"
