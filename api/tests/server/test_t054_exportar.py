"""T-054 — `/api/componentes/{id}/exportar` contra o app real. AC-1, AC-4..AC-9.

A exportacao e' a leitura com outra embalagem. O risco e' ela virar o atalho:
uma checagem que `dados` faz e `exportar` esqueceu, e o arquivo tem o que a
tela recusaria. Por isso os negativos daqui sao os mesmos de `dados`, rodados
contra as duas rotas, e o positivo de cada um prova que a rota nao recusa tudo.

Pula sem banco: `make db-local && make db-teste`.
"""

import csv
import io
import secrets
from collections.abc import AsyncGenerator
from typing import Any

import pytest
import sqlalchemy as sa
from borda import cliente, criar_sessao, descartar_pool, motor_dono
from openpyxl import load_workbook
from pypdf import PdfReader

from estoque.server.rotas import exportar as rota_exportar


@pytest.fixture(autouse=True)
async def _pool_limpo() -> AsyncGenerator[None]:
    yield
    await descartar_pool()


RELATORIO: dict[str, object] = {"agrupar_por": "unidade", "periodo": "365"}


def _sessao(papel: str, unidades: tuple[str, ...] = ("cd-matriz",)) -> tuple[str, str]:
    uid = f"t054-{papel}-{secrets.token_hex(4)}"
    return uid, criar_sessao(usuario_id=uid, papel=papel, unidades=unidades)


async def _exportar(
    sid: str | None,
    componente: str,
    params: dict[str, object],
    formato: str = "csv",
    **kw: Any,
) -> Any:
    async with cliente(sid, **kw) as cli:
        return await cli.post(
            f"/api/componentes/{componente}/exportar",
            json={"params": params, "formato": formato},
        )


async def _dados(sid: str, componente: str, params: dict[str, object]) -> Any:
    async with cliente(sid) as cli:
        return await cli.post(
            f"/api/componentes/{componente}/dados",
            json={"params": params, "pagina": {"limite": 200}},
        )


def _tabela_csv(corpo: bytes) -> list[list[str]]:
    linhas = list(csv.reader(io.StringIO(corpo.decode("utf-8-sig")), delimiter=";"))
    inicio = next(i for i, x in enumerate(linhas) if not x) + 1
    return linhas[inicio:]


def _texto_xlsx(corpo: bytes) -> str:
    ws = load_workbook(io.BytesIO(corpo)).worksheets[0]
    return "\n".join(str(c.value) for r in ws.iter_rows() for c in r if c.value is not None)


def _texto_pdf(corpo: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(corpo)).pages)


def _exportacoes(uid: str) -> list[dict[str, Any]]:
    with motor_dono().connect() as c:
        return [
            dict(r._mapping)
            for r in c.execute(
                sa.text(
                    "SELECT entidade, valor_novo, origem FROM auditoria "
                    "WHERE acao='exportar' AND ator_id=:u ORDER BY id"
                ),
                {"u": uid},
            )
        ]


# --- AC-1 · os tres formatos, com as linhas da tela --------------------------


async def test_ac1_csv_do_relatorio_tem_as_mesmas_linhas_de_dados() -> None:
    _, sid = _sessao("diretor")
    tela = (await _dados(sid, "relatorio_movimentacao", RELATORIO)).json()["dados"]
    r = await _exportar(sid, "relatorio_movimentacao", RELATORIO)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("text/csv")
    assert r.headers["content-disposition"].startswith(
        'attachment; filename="movimentacao-por-unidade-'
    )
    assert r.headers["cache-control"] == "no-store"
    corpo = _tabela_csv(r.content)
    assert [x[0] for x in corpo[1:-1]] == [x["rotulo"] for x in tela["linhas"]]
    # Inteiro sai sem separador de milhar: e' o que o Excel le como numero.
    assert [int(x[1]) for x in corpo[1:-1]] == [x["numero"] for x in tela["linhas"]]
    assert corpo[-1][:2] == ["Total", str(tela["total"])]


@pytest.mark.parametrize(
    ("formato", "midia", "ler"),
    [
        ("xlsx", "application/vnd.openxmlformats", _texto_xlsx),
        ("pdf", "application/pdf", _texto_pdf),
    ],
)
async def test_ac1_xlsx_e_pdf_do_relatorio(formato: str, midia: str, ler: Any) -> None:
    _, sid = _sessao("diretor")
    tela = (await _dados(sid, "relatorio_movimentacao", RELATORIO)).json()["dados"]
    r = await _exportar(sid, "relatorio_movimentacao", RELATORIO, formato)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(midia)
    assert r.headers["content-disposition"].endswith(f'.{formato}"')
    texto = ler(r.content)
    for linha in tela["linhas"]:
        assert linha["rotulo"] in texto


@pytest.mark.parametrize(
    ("componente", "params", "papel", "unidades"),
    [
        (
            "temperatura_historico",
            {"unidade_id": "cd-refrigerado", "de": "2026-01-01", "ate": "2026-12-31"},
            "rt",
            ("cd-refrigerado",),
        ),
        (
            "temperatura_excursoes",
            {"unidade_id": "cd-refrigerado", "de": "2026-01-01", "ate": "2026-12-31"},
            "rt",
            ("cd-refrigerado",),
        ),
        ("fila_vencimento", {}, "gerente", ("cd-matriz",)),
        ("lote_lista", {}, "gerente", ("cd-matriz",)),
        ("movimento_lista", {}, "rt", ("cd-matriz",)),
        ("auditoria_trilha", {}, "auditoria", ("cd-matriz",)),
    ],
)
@pytest.mark.parametrize("formato", ["csv", "xlsx", "pdf"])
async def test_ac1_os_sete_exportaveis_geram_os_tres_formatos(
    componente: str,
    params: dict[str, object],
    papel: str,
    unidades: tuple[str, ...],
    formato: str,
) -> None:
    _, sid = _sessao(papel, unidades)
    assert (await _dados(sid, componente, params)).status_code == 200
    r = await _exportar(sid, componente, params, formato)
    assert r.status_code == 200, r.text
    assert len(r.content) > 0


# --- AC-4 · auditoria --------------------------------------------------------


async def test_ac4_exportar_grava_auditoria_com_formato_e_params() -> None:
    uid, sid = _sessao("diretor")
    r = await _exportar(sid, "relatorio_movimentacao", RELATORIO, "xlsx")
    assert r.status_code == 200
    (linha,) = _exportacoes(uid)
    assert linha["entidade"] == "relatorio_movimentacao"
    assert linha["origem"] == "tela"
    assert linha["valor_novo"]["formato"] == "xlsx"
    assert linha["valor_novo"]["params"] == RELATORIO
    assert linha["valor_novo"]["linhas"] >= 1


async def test_ac4_exportacao_recusada_nao_registra_arquivo_entregue() -> None:
    """Negativo: o registro diz que um arquivo SAIU. Recusa nao entrega nada,
    e registrar `exportar` para ela seria trilha dizendo o que nao aconteceu."""
    uid, sid = _sessao("rt")
    r = await _exportar(sid, "relatorio_movimentacao", {**RELATORIO, "metrica": "valor"})
    assert r.status_code == 403
    assert _exportacoes(uid) == []


# --- AC-5 · custo nao vaza (RF-14, RN-A02) -----------------------------------


@pytest.mark.parametrize("formato", ["csv", "xlsx", "pdf"])
async def test_ac5_rt_pedindo_valor_e_recusado_sem_arquivo(formato: str) -> None:
    _, sid = _sessao("rt")
    r = await _exportar(
        sid, "relatorio_movimentacao", {**RELATORIO, "metrica": "valor"}, formato
    )
    assert r.status_code == 403, r.text
    assert r.headers["content-type"].startswith("application/json")
    assert "nao_autorizado" in r.text


@pytest.mark.parametrize(
    ("formato", "ler"),
    [("csv", lambda b: b.decode("utf-8-sig")), ("xlsx", _texto_xlsx), ("pdf", _texto_pdf)],
)
async def test_ac5_arquivo_do_rt_nao_tem_dinheiro(formato: str, ler: Any) -> None:
    _, sid = _sessao("rt")
    for metrica in ("quantidade", "movimentos"):
        r = await _exportar(
            sid,
            "relatorio_movimentacao",
            {**RELATORIO, "agrupar_por": "produto", "metrica": metrica},
            formato,
        )
        assert r.status_code == 200, r.text
        texto = ler(r.content)
        # Nao "custo": o CS-04 cadastra produto com essa palavra no nome.
        assert "R$" not in texto
        assert "Valor" not in texto


async def test_ac5_contraponto_o_diretor_exporta_valor_em_reais() -> None:
    _, sid = _sessao("diretor")
    r = await _exportar(
        sid, "relatorio_movimentacao", {**RELATORIO, "metrica": "valor"}, "xlsx"
    )
    assert r.status_code == 200, r.text
    assert "Valor (R$)" in _texto_xlsx(r.content)


# --- AC-6 · a rota nao e' atalho (os mesmos negativos de `dados`) ------------

NEGATIVOS = [
    # Rafael (comprador) nao tem `movimento.ler`.
    pytest.param(
        "comprador", ("cd-matriz",), "movimento_lista", {}, 403, id="fora-do-catalogo"
    ),
    # Odair alcanca so' Uberlandia; pedir a Matriz nega (ADR-0014).
    pytest.param(
        "gerente",
        ("filial-uberlandia",),
        "lote_lista",
        {"unidade_id": "cd-matriz"},
        403,
        id="unidade-fora-do-escopo",
    ),
    pytest.param(
        "rt",
        ("cd-matriz",),
        "relatorio_movimentacao",
        {**RELATORIO, "metrica": "valor"},
        403,
        id="valor-sem-custo",
    ),
    pytest.param("conferente", ("cd-matriz",), "auditoria_trilha", {}, 403, id="trilha"),
]


@pytest.mark.parametrize("rota", ["dados", "exportar"])
@pytest.mark.parametrize(("papel", "unidades", "componente", "params", "codigo"), NEGATIVOS)
async def test_ac6_mesma_recusa_em_dados_e_em_exportar(
    rota: str,
    papel: str,
    unidades: tuple[str, ...],
    componente: str,
    params: dict[str, object],
    codigo: int,
) -> None:
    _, sid = _sessao(papel, unidades)
    if rota == "dados":
        r = await _dados(sid, componente, params)
    else:
        r = await _exportar(sid, componente, params)
    assert r.status_code == codigo, r.text


async def test_ac6_contraponto_o_gerente_exporta_a_propria_unidade() -> None:
    _, sid = _sessao("gerente", ("filial-uberlandia",))
    r = await _exportar(sid, "lote_lista", {"unidade_id": "filial-uberlandia"})
    assert r.status_code == 200, r.text


async def test_ac6_sem_sessao_e_recusado() -> None:
    motor_dono()  # pula sem banco, como os outros: a rota consulta a sessao
    r = await _exportar(None, "lote_lista", {})
    assert r.status_code == 401


async def test_ac6_sem_csrf_e_recusado() -> None:
    _, sid = _sessao("gerente")
    r = await _exportar(sid, "lote_lista", {}, com_csrf=False)
    assert r.status_code == 403


async def test_ac6_formato_fora_do_enum_e_recusado() -> None:
    _, sid = _sessao("gerente")
    r = await _exportar(sid, "lote_lista", {}, "docx")
    assert r.status_code == 422


# --- AC-8 · teto de linhas: recusa, nunca corte ------------------------------


async def test_ac8_acima_do_teto_recusa_e_manda_estreitar(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _, sid = _sessao("gerente")
    monkeypatch.setattr(rota_exportar, "TETO_DE_LINHAS", 1)
    r = await _exportar(sid, "lote_lista", {})
    assert r.status_code == 422, r.text
    assert "Estreite o filtro" in r.json()["erro"]["mensagem"]


async def test_ac8_lista_sem_paginacao_tambem_respeita_o_teto(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`movimento_lista` nao pagina: o teto e' conferido na tabela, nao so' no
    `tem_mais`."""
    _, sid = _sessao("rt")
    monkeypatch.setattr(rota_exportar, "TETO_DE_LINHAS", 1)
    r = await _exportar(sid, "movimento_lista", {})
    assert r.status_code == 422, r.text


async def test_ac8_contraponto_no_teto_real_exporta() -> None:
    _, sid = _sessao("gerente")
    r = await _exportar(sid, "lote_lista", {})
    assert r.status_code == 200, r.text


# --- AC-9 · componente nao exportavel ----------------------------------------


async def test_ac9_componente_fora_do_registro_de_exportacao_e_nao_encontrado() -> None:
    _, sid = _sessao("diretor")
    assert (
        await _exportar(sid, "estoque_indicador", {"metrica": "lotes_ativos"})
    ).status_code == 404
    assert (await _exportar(sid, "nao_existe", {})).status_code == 404


async def test_ac9_catalogo_diz_quem_exporta() -> None:
    _, sid = _sessao("diretor")
    async with cliente(sid) as cli:
        cat = (await cli.get("/api/catalogo")).json()["dados"]
    marcado = {e["id"]: e["exportavel"] for e in cat}
    assert marcado["relatorio_movimentacao"] is True
    assert marcado["estoque_indicador"] is False
